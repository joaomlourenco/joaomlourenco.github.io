#!/usr/bin/env python3
"""Worksheet to collect the LinkedIn profiles of the advisees that do not have one yet.

    python3 scripts/linkedin_worksheet.py [--degree BSc] [-o ~/Desktop/linkedin-worksheet.xlsx]

One row per advisee page without a `linkedin:` front-matter field: year, name, company (from the
internship/co-adviser fields), two LinkedIn *search* links (nothing is looked up by this script --
you open the links, recognise the right person and paste the profile URL in the "LinkedIn URL"
column), and the other advisee pages with the same first+last name (to disambiguate).
Afterwards:  python3 scripts/linkedin_apply.py <the sheet>   (dry run; --write to apply).

Writes an .xlsx (needs openpyxl) or, with a .csv name, a CSV.
"""
import argparse, csv, glob, os, re, sys, urllib.parse

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ADV = os.path.join(ROOT, "content", "advisees")
SUFFIX = re.compile(r"[,\s]+(s\.?\s?a\.?|lda\.?|ltda\.?|unipessoal|lda)\s*$", re.I)


def front_matter(path):
    t = open(path, encoding="utf-8").read()
    m = re.match(r"---\n(.*?)\n---", t, re.S)
    return m.group(1) if m else ""


def field(fm, key):
    m = re.search(r"^\s*" + key + r':\s*"?(.*?)"?\s*$', fm, re.M)
    return m.group(1) if m else ""


def company(fm):
    co = field(fm, "coadviser") or field(fm, "mainadviser")
    m = re.search(r"\(([^()]*(?:\([^()]*\))?[^()]*)\)\s*$", co)
    c = m.group(1) if m else ""
    if not c:
        o = field(fm, "other")
        m = re.search(r"[Ii]nternship at (.+?)(?:,\s*Lisboa.*|\.)?$", o)
        c = m.group(1) if m else ("NOVA FCT" if re.search(r"DI-FCT|FCT-NOVA|research|project", o, re.I) else "")
    c = re.split(r"\s[-–—]\s|,\s", c)[0]
    return SUFFIX.sub("", c).strip()


def short_name(name):
    w = name.split()
    return w[0] + " " + w[-1] if len(w) > 2 else name


def search(*terms):
    return "https://www.linkedin.com/search/results/people/?keywords=" + urllib.parse.quote(" ".join(t for t in terms if t))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--degree", default="BSc")
    ap.add_argument("-o", "--output", default=os.path.expanduser("~/Desktop/linkedin-worksheet-%s.xlsx" % "BSc"))
    a = ap.parse_args()
    pages = []
    for f in sorted(glob.glob(os.path.join(ADV, "*", "*", "index.md"))):
        fm = front_matter(f)
        pages.append((f, fm, field(fm, "title")))
    byname = {}
    for f, fm, t in pages:
        byname.setdefault(short_name(t).lower(), []).append(f)
    rows = []
    for f, fm, t in pages:
        if field(fm, "degree") != a.degree or "linkedin" in fm:
            continue
        rel = os.path.relpath(f, ROOT)
        same = [os.path.relpath(g, ADV).split(os.sep)[0] + " " + field(front_matter(g), "degree")
                for g in byname[short_name(t).lower()] if g != f]
        co = company(fm)
        rows.append([os.path.relpath(f, ADV).split(os.sep)[0], t, co,
                     search(short_name(t), co), search(short_name(t), "NOVA FCT Informática"), "",
                     "; ".join(same), rel])
    rows.sort(key=lambda r: (-int(r[0]), r[1]))
    head = ["Year", "Name", "Company", "Search: name + company", "Search: name + NOVA FCT", "LinkedIn URL (paste here)",
            "Other advisee pages with this first+last name", "Page"]
    if a.output.endswith(".csv"):
        with open(a.output, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh); w.writerow(head); w.writerows(rows)
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "LinkedIn"
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
            n = ws.max_row
            for col, label in ((4, "search"), (5, "search")):
                cell = ws.cell(row=n, column=col); cell.hyperlink = r[col - 1]; cell.value = label
                cell.font = Font(color="0563C1", underline="single")
            ws.cell(row=n, column=6).fill = PatternFill("solid", fgColor="FFF2CC")
            if r[6]:
                ws.cell(row=n, column=7).font = Font(color="C00000", bold=True)
        for col, wd in zip("ABCDEFGH", (7, 38, 34, 12, 12, 52, 38, 50)):
            ws.column_dimensions[col].width = wd
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(a.output)
    print("%d rows -> %s" % (len(rows), a.output))


if __name__ == "__main__":
    main()
